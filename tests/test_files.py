import io

from docx import Document

from files import translate_file, translate_srt

SRT = """1
00:00:01,000 --> 00:00:02,000
Hello there

2
00:00:03,000 --> 00:00:04,000
Line one
Line two
"""


def test_srt_keeps_numbers_and_timestamps():
    out = translate_srt(SRT, str.upper)
    assert "00:00:01,000 --> 00:00:02,000" in out
    assert "HELLO THERE" in out and "LINE ONE\nLINE TWO" in out
    assert out.startswith("1\n")


def test_txt_and_names():
    name, data = translate_file("notes.txt", "hi".encode(), str.upper)
    assert name == "notes.translated.txt" and data == b"HI"


def test_docx_translates_paragraphs():
    doc = Document()
    doc.add_paragraph("hello world")
    buf = io.BytesIO()
    doc.save(buf)
    name, data = translate_file("a.docx", buf.getvalue(), str.upper)
    assert name == "a.translated.docx"
    assert Document(io.BytesIO(data)).paragraphs[0].text == "HELLO WORLD"


def test_unsupported():
    import pytest

    with pytest.raises(ValueError):
        translate_file("a.pdf", b"", str)


def _docx_bytes(build):
    doc = Document()
    build(doc)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def test_docx_keeps_inline_formatting():
    def build(doc):
        p = doc.add_paragraph()
        p.add_run("Say ")
        p.add_run("hello").bold = True
        p.add_run(" to everyone")

    _, data = translate_file("a.docx", _docx_bytes(build), lambda s: "Dis bonjour a tout le monde")
    runs = Document(io.BytesIO(data)).paragraphs[0].runs
    assert "".join(r.text for r in runs) == "Dis bonjour a tout le monde"
    bold_text = "".join(r.text for r in runs if r.bold)
    assert bold_text.strip() and bold_text.strip() in "Dis bonjour a tout le monde"
    assert not runs[0].bold and not runs[-1].bold


def test_docx_hyperlink_text_is_translated_not_duplicated():
    from docx.oxml import parse_xml
    from docx.oxml.ns import nsdecls

    def build(doc):
        p = doc.add_paragraph()
        p.add_run("See ")
        p._p.append(parse_xml(
            f'<w:hyperlink {nsdecls("w")} w:anchor="x"><w:r><w:t>this page</w:t></w:r></w:hyperlink>'
        ))

    _, data = translate_file("a.docx", _docx_bytes(build), str.upper)
    assert Document(io.BytesIO(data)).paragraphs[0].text == "SEE THIS PAGE"


def test_docx_parallel_keeps_order_and_tables():
    def build(doc):
        for word in ("one", "two", "three", "four"):
            doc.add_paragraph(word)
        table = doc.add_table(rows=1, cols=1)
        table.cell(0, 0).text = "cell"

    _, data = translate_file("a.docx", _docx_bytes(build), str.upper, workers=4)
    out = Document(io.BytesIO(data))
    assert [p.text for p in out.paragraphs if p.text] == ["ONE", "TWO", "THREE", "FOUR"]
    assert out.tables[0].cell(0, 0).text == "CELL"


def test_spread_over_runs_is_lossless():
    from files import spread_over_runs

    for translated in ("", "x", "Dis bonjour a tout le monde", "日本語のテキスト"):
        pieces = spread_over_runs(["Say ", "hello", " to everyone"], ["a", "b", "a"], translated)
        assert "".join(pieces) == translated and len(pieces) == 3
